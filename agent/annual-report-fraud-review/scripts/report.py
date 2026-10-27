"""Export a readable report and the exact machine-readable review records."""
import io
import re
import zipfile
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from workpaper import STATUS, coverage_counts, digest, knowledge, write


def export(w, calculations, validation, output):
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    values = {}
    for f in w["facts"]:
        values[f["id"]] = (f["value"], f["context"]["unit"]) if f["verified"] else (None, f["context"]["unit"])
    for c in calculations["calculations"]:
        value = c["value"]
        if value is not None:
            value = format(Decimal(value), ".6f").rstrip("0").rstrip(".")
            if value in ("", "-0"):
                value = "0"
        values[c["id"]] = value, c["context"]["unit"]

    def resolve(text):
        def substitute(match):
            key = match.group(1)
            value, unit = values[key]
            return f"{value}{unit}〔{key}〕" if value is not None else f"未确定〔{key}〕"
        return re.sub(r"\{\{([^{}]+)\}\}", substitute, str(text))

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.7)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(10)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.2
    doc.core_properties.title = "企业年报财务舞弊风险排查报告"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = ""
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    markdown = []

    def heading(text, level=1):
        text = resolve(text)
        doc.add_heading(text, level=level)
        markdown.append("#" * (level + 1) + " " + text + "\n")

    def paragraph(text, interpolate=True):
        text = resolve(text) if interpolate else str(text)
        doc.add_paragraph(text)
        markdown.append(text + "\n")

    def table(headers, rows, widths):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.autofit = False
        for column, width in zip(t.columns, widths):
            column.width = Inches(width)
        for cell, header, width in zip(t.rows[0].cells, headers, widths):
            cell.text = header
            cell.width = Inches(width)
            for run in cell.paragraphs[0].runs:
                run.bold = True
        repeat = OxmlElement("w:tblHeader")
        t.rows[0]._tr.get_or_add_trPr().append(repeat)
        escaped = lambda x: resolve(x).replace("|", "\\|").replace("\n", "<br>")
        markdown.extend(["| " + " | ".join(headers) + " |", "| " + " | ".join("---" for _ in headers) + " |"])
        for row in rows:
            cells = t.add_row().cells
            for cell, value, width in zip(cells, row, widths):
                cell.text = resolve(value)
                cell.width = Inches(width)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(3)
                    for run in p.runs:
                        run.font.size = Pt(9)
            markdown.append("| " + " | ".join(escaped(v) for v in row) + " |")
        markdown.append("")

    counts = coverage_counts(w)
    unfinished = counts["not_completed"] or any(p["status"] != "reviewed" for p in w["pages"]) or not w["sections"] or any(
        s["status"] != "reviewed" for s in w["sections"]) or any(f["review"]["status"] != "verified" for f in w["findings"]) or not w["summary"].strip() or any(
        w["report"][k] is None for k in ("company", "year", "scope", "accounting_basis"))
    progress = "部分完成 存在未完成程序或原页复核" if unfinished else "本范围程序已登记完成 资料缺口仍按条目保留"
    doc.add_heading("企业年报财务舞弊风险排查报告", level=0)
    markdown.append("# 企业年报财务舞弊风险排查报告\n")
    paragraph(f"企业：{w['report']['company'] or '未确定'}　年度：{w['report']['year'] or '未确定'}")
    paragraph("进度：" + progress)
    paragraph("本报告仅以指定的一份年报为公司事实来源，覆盖其实际披露的比较数据。结论用于财务舞弊风险排查，不提供舞弊概率或审计鉴证意见。")
    paragraph("分析范围：" + str(w["report"]["scope"] or "未确定") + "；会计基础：" + str(w["report"]["accounting_basis"] or "未确定"))
    heading("重点结论")
    paragraph(w["summary"] or "尚未形成摘要，以下为当前已保存的证据及程序进度。")
    heading("覆盖情况和资料限制")
    paragraph("逐项登记不等于完成检查。已执行检查、不适用、证据不足和未完成分别列示，不合并为一个通过率。")
    table(["状态", "条目数"], [(STATUS[s], str(counts[s])) for s in STATUS], [5.5, 1.37])
    for warning in validation["warnings"]:
        paragraph(warning)
    for limitation in w["limitations"]:
        paragraph(limitation)
    heading("重要发现及反证检验")
    if not w["findings"]:
        paragraph("当前未登记发现；这不代表全部程序已执行或不存在异常。")
    priorities = {"high": "高", "medium": "中", "low": "低"}
    strengths = {"strong": "较强", "moderate": "中等", "limited": "有限"}
    for f in sorted(w["findings"], key=lambda x: ("high", "medium", "low").index(x["priority"])):
        heading(f["id"] + " " + f["title"], 2)
        paragraph(f"核查优先级：{priorities[f['priority']]}；证据强度：{strengths[f['evidence_strength']]}；规则：{', '.join(f['rule_ids']) or '清单外发现'}")
        for label, key in (("已披露事实", "observations"), ("可能机制", "mechanism"), ("当前判断", "assessment"), ("可能影响", "impact")):
            paragraph(label + "：" + f[key])
        paragraph("支持证据：" + ", ".join(f["evidence_ids"]) + "；反证引用：" + (", ".join(f["counter_evidence_ids"]) or "未取得，不能视为不存在"))
        paragraph("复算引用：" + (", ".join(f["calculation_ids"]) or "无已登记计算"))
        for explanation in f["alternative_explanations"]:
            paragraph("替代解释及核查：" + explanation)
        for missing in f["remaining_evidence"]:
            paragraph("尚缺资料：" + missing)
        paragraph("改变或关闭判断的条件：" + f["close_when"])
        if f["overlap_group"]:
            paragraph("共同影响组：" + f["overlap_group"] + "。与同组发现不可重复累计。")
        paragraph("原页复核：" + ("已登记复核" if f["review"]["status"] == "verified" else "待复核") + "；" + f["review"]["notes"])
    heading("专题排查全项清单")
    bundle = knowledge()
    by_id = {r["id"]: r for r in bundle["rules"]}
    coverage = {c["rule_id"]: c for c in w["coverage"]}
    for topic in bundle["topics"]:
        heading(topic["title"], 2)
        rows = []
        for rule_id in topic["rule_ids"]:
            c = coverage[rule_id]
            detail = c["reason"] or "尚未执行"
            if c["procedures"]:
                detail += "\n实际程序：" + "；".join(c["procedures"])
            if c["evidence_ids"]:
                detail += "\n证据：" + ", ".join(c["evidence_ids"])
            if c["calculation_ids"] or c["finding_ids"]:
                detail += "\n计算/发现：" + ", ".join(c["calculation_ids"] + c["finding_ids"])
            if c["limitations"]:
                detail += "\n缺口：" + "；".join(c["limitations"])
            rows.append((rule_id, by_id[rule_id]["method"], STATUS[c["status"]], detail))
        table(["ID", "排查事项", "状态", "依据和实际程序"], rows, [0.45, 1.65, 1.05, 3.72])
    heading("阅读范围")
    for s in w["sections"]:
        paragraph(f"{s['title']}；物理页：{', '.join(map(str, s['pages']))}；状态：{s['status']}；{s['notes']}")
    unread = [p for p in w["pages"] if p["status"] != "reviewed"]
    for p in unread:
        paragraph(f"物理页 {p['page']}：{p['status']}；{p['notes']}")
    heading("证据索引")
    for e in w["evidence"]:
        paragraph(f"{e['id']}；PDF 物理页 {e['page']}；{e['locator']}；核实状态：{e['verification']}", False)
        paragraph(e["quote"], False)
        if e["table_header"]:
            paragraph("表头：" + " / ".join(e["table_header"]) + "；行：" + (e["row"] or "未登记") + "；列：" + (e["column"] or "未登记"), False)
    heading("事实和复算底稿")
    for f in w["facts"]:
        ctx = f["context"]
        paragraph(f"{f['id']} {f['label']}：{f['value'] if f['value'] is not None else '未披露'}{ctx['unit']}；{ctx['entity']} / {ctx['scope']} / {ctx['basis']}；{ctx['start'] or '时点'} 至 {ctx['end']}；币种：{ctx['currency'] or '无'}；核实：{f['verified']}；证据：{', '.join(f['evidence_ids'])}", False)
    paragraph("计算展示保留至 6 位小数，完整十进制结果保存在 calculations.json。")
    for c in calculations["calculations"]:
        terms = ", ".join(f"{t['ref']} × {t['weight']}" for t in c["terms"])
        paragraph(f"{c['id']} {c['label']}：" + "{{" + c["id"] + "}}" + f"；运算：{c['op']}({terms}) × {c['multiplier']}；{c['reason'] or '已计算'}")
        paragraph("测算依据：" + c["rationale"])
        if c["period_rule"] == "sensitivity":
            paragraph("本项为敏感性测算，表示所述比例条件下的影响，不代表应确认的会计金额。")
    for r in calculations["reconciliations"]:
        paragraph(f"{r['id']}：{r['actual']} − {r['expected']}；残差：{r['residual'] if r['residual'] is not None else '未确定'} {r['unit'] or ''}；状态：{r['status']}；容差：{r['tolerance']}；依据：{r['basis']}")
    heading("版本与来源")
    paragraph(f"输入文件：{w['report']['filename']}；物理页数：{w['report']['page_count']}；SHA256：{w['report']['sha256']}", False)
    paragraph(f"知识版本：{w['knowledge']['version']}；规则 SHA256：{w['knowledge']['rules_sha256']}；图谱 SHA256：{w['knowledge']['graph_sha256']}")
    write(output / "workpaper.json", w)
    calculations["workpaper_sha256"] = digest(output / "workpaper.json")
    write(output / "calculations.json", calculations)
    (output / "report.md").write_text("\n".join(markdown) + "\n", encoding="utf-8")
    buffer = io.BytesIO()
    doc.save(buffer)
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(output / "report.docx", "w", zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == "docProps/app.xml":
                tree = ET.fromstring(content)
                for item in list(tree):
                    if item.tag.rsplit("}", 1)[-1] in ("Application", "AppVersion"):
                        tree.remove(item)
                content = ET.tostring(tree, encoding="utf-8", xml_declaration=True)
            target.writestr(entry, content)
    return {"files": ["workpaper.json", "calculations.json", "report.md", "report.docx"],
            "progress": progress, "coverage": counts, "warnings": validation["warnings"]}
